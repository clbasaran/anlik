#!/usr/bin/env python3
"""Excel dosyasından SeedData.swift üretir.

Kullanım:
    pip install openpyxl
    python3 scripts/generate_seed.py docs/soulmate_coffee_stok_takip.xlsx

STOK TAKİBİ sayfasındaki ürün kataloğunu ve ZAYİ TAKİBİ sayfasındaki menü
listesini okuyup SoulmateCoffee/Store/SeedData.swift dosyasını yeniden yazar.
"""
import sys
import os
import openpyxl

GROUP_MAP = {
    "SICAK İÇECEK": ".sicakIcecek",
    "SOĞUK İÇECEK": ".sogukIcecek",
    "GIDA": ".gida",
    "SARF": ".sarf",
    "TEMİZLİK": ".temizlik",
}


def esc(s: str) -> str:
    return s.replace("\\", "\\\\").replace('"', '\\"')


def main(xlsx_path: str) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["STOK TAKİBİ"]

    stock_lines = []
    for r in range(4, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if name is None:
            continue
        name = str(name).strip()
        if name == "TOPLAM":
            continue
        group = GROUP_MAP[str(ws.cell(r, 2).value).strip()]
        unit = str(ws.cell(r, 3).value).strip()
        price = ws.cell(r, 4).value or 0
        price = int(price) if float(price).is_integer() else price
        stock_lines.append(
            f'        StockItem(name: "{esc(name)}", group: {group}, '
            f'unit: "{esc(unit)}", unitPrice: {price}),'
        )

    wz = wb["ZAYİ TAKİBİ"]
    menu = []
    for r in range(2, wz.max_row + 1):
        a = wz.cell(r, 1).value
        if a is None or not str(a).strip():
            continue
        s = str(a).strip()
        if s.upper() == "ÜRÜN ADI":
            continue
        menu.append(s)
    menu_lines = [f'        "{esc(m)}",' for m in menu]

    out = '''import Foundation

/// Excel dosyasından (Soulmate Coffee - Kötekli Şubesi, 2026 Ocak) birebir
/// aktarılan başlangıç kataloğu. Yeni bir dönem oluşturulduğunda bu liste
/// kullanılır. Bu dosya scripts/generate_seed.py ile xlsx'ten üretilmiştir.
enum SeedData {

    static let branchName = "SOULMATE COFFEE - KÖTEKLİ ŞUBESİ"

    /// STOK TAKİBİ ürün kataloğu.
    static let stockCatalog: [StockItem] = [
%s
    ]

    /// ZAYİ / İKRAM günlük takip ürün listesi (menü kalemleri).
    static let dailyCatalog: [String] = [
%s
    ]

    /// Belirli bir dönem için boş (giriş bekleyen) AppData üretir.
    static func makePeriod(year: Int, month: Int) -> AppData {
        AppData(
            branchName: branchName,
            year: year,
            month: month,
            ciro: 0,
            calisanGunSayisi: 0,
            personelMaliyet: 0,
            kira: 0,
            elektrik: 0,
            su: 0,
            dogalgaz: 0,
            internet: 0,
            digerGiderler: 0,
            stock: stockCatalog,
            zayi: dailyCatalog.map { DailyTrackItem(name: $0) },
            ikram: dailyCatalog.map { DailyTrackItem(name: $0) }
        )
    }
}
''' % ("\n".join(stock_lines), "\n".join(menu_lines))

    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    dest = os.path.join(here, "SoulmateCoffee", "Store", "SeedData.swift")
    with open(dest, "w", encoding="utf-8") as f:
        f.write(out)
    print(f"Yazıldı: {dest}  ({len(stock_lines)} ürün, {len(menu_lines)} menü kalemi)")


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else "docs/soulmate_coffee_stok_takip.xlsx"
    main(path)
